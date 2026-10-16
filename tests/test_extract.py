import io

import pytest

from federal_ja.collect import extract as ex


def text_pdf(pages: list[str]) -> bytes:
    """Minimal hand-built PDF with one Helvetica text line per page."""
    objs = []
    n = len(pages)
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(n))
    objs.append("<< /Type /Catalog /Pages 2 0 R >>")
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {n} >>")
    font_id = 3 + 2 * n
    for i, t in enumerate(pages):
        content = f"BT /F1 10 Tf 20 700 Td ({t}) Tj ET".encode()
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                    f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {4 + 2 * i} 0 R >>")
        objs.append(content)
    objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, o in enumerate(objs, 1):
        offsets.append(out.tell())
        if isinstance(o, bytes):
            out.write(f"{i} 0 obj\n<< /Length {len(o)} >>\nstream\n".encode() + o + b"\nendstream\nendobj\n")
        else:
            out.write(f"{i} 0 obj\n{o}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def blank_pdf(n: int) -> bytes:
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(n):
        w.add_blank_page(612, 792)
    b = io.BytesIO()
    w.write(b)
    return b.getvalue()


LONG = "Justification for other than full and open competition " * 4  # >100 alnum chars


def test_classification_rule_is_per_page():
    assert ex.pdf_needs_transcription([]) is True
    assert ex.pdf_needs_transcription([""]) is True
    assert ex.pdf_needs_transcription([LONG]) is False
    assert ex.pdf_needs_transcription([LONG, LONG]) is False
    # one rasterized page in an otherwise typed document still needs transcription
    assert ex.pdf_needs_transcription([LONG, LONG, ""]) is True
    assert ex.textless_pages([LONG, "", LONG, "Page 4 of 4"]) == [2, 4]


def test_threshold_boundary():
    just_under = "a" * (ex.MIN_ALNUM_PER_TEXT_PAGE - 1)
    at = "a" * ex.MIN_ALNUM_PER_TEXT_PAGE
    assert ex.textless_pages([just_under, at]) == [1]
    # punctuation and whitespace do not count
    assert ex.textless_pages(["- . , ;" * 100]) == [1]


def test_text_pdf(tmp_path):
    p = tmp_path / "a.pdf"
    p.write_bytes(text_pdf([LONG, LONG]))
    r = ex.extract_file(str(p), "a.pdf")
    assert r.file_kind == "pdf" and r.page_count == 2 and r.pages_with_text == 2
    assert r.text_source == "pdf_text" and not r.needs_transcription and r.textless_pages == []
    assert "--- page 1 ---" in r.text and "--- page 2 ---" in r.text
    assert "Justification for other" in r.text


def test_scanned_pdf(tmp_path):
    p = tmp_path / "s.pdf"
    p.write_bytes(blank_pdf(3))
    r = ex.extract_file(str(p), "s.pdf")
    assert r.needs_transcription and r.text_source == "none" and r.text == ""
    assert r.page_count == 3 and r.pages_with_text == 0 and r.textless_pages == [1, 2, 3]


def test_mixed_pdf_keeps_native_text_and_flags_pages(tmp_path):
    p = tmp_path / "m.pdf"
    p.write_bytes(text_pdf([LONG, "x", LONG]))
    r = ex.extract_file(str(p), "m.pdf")
    assert r.needs_transcription and r.textless_pages == [2] and r.pages_with_text == 2
    assert r.text_source == "pdf_text" and "--- page 3 ---" in r.text


def test_corrupt_pdf_records_error(tmp_path):
    p = tmp_path / "c.pdf"
    p.write_bytes(b"%PDF-1.4\ngarbage")
    r = ex.extract_file(str(p), "c.pdf")
    assert r.file_kind == "pdf" and r.text_source == "none"
    assert r.extract_error or r.page_count == 0


def test_docx(tmp_path):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_paragraph("Justification and Approval")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Amount", "$1,000"
    p = tmp_path / "a.docx"
    d.save(str(p))
    r = ex.extract_file(str(p), "A.DOCX")
    assert r.file_kind == "docx" and r.text_source == "docx"
    assert "Justification and Approval" in r.text and "Amount\t$1,000" in r.text


def test_docx_detected_without_extension(tmp_path):
    docx = pytest.importorskip("docx")
    d = docx.Document()
    d.add_paragraph("hello")
    p = tmp_path / "noext"
    d.save(str(p))
    assert ex.extract_file(str(p), "noext").text_source == "docx"


def test_xlsx(tmp_path):
    import openpyxl
    wb = openpyxl.Workbook()
    wb.active.title = "Costs"
    wb.active.append(["Item", "Price"])
    wb.active.append(["Widget", 5])
    p = tmp_path / "a.xlsx"
    wb.save(str(p))
    r = ex.extract_file(str(p), "a.xlsx")
    assert r.text_source == "xlsx"
    assert "=== sheet: Costs ===" in r.text and "Widget\t5" in r.text


def test_unknown_and_legacy(tmp_path):
    p = tmp_path / "a.doc"
    p.write_bytes(b"\xd0\xcf\x11\xe0" + b"\0" * 100)
    r = ex.extract_file(str(p), "a.doc")
    assert r.file_kind == "ole" and r.text_source == "none"
    p2 = tmp_path / "x.bin"
    p2.write_bytes(b"\x00\x01")
    assert ex.extract_file(str(p2), "x.bin").file_kind == "other"


def test_text_file_latin1(tmp_path):
    p = tmp_path / "a.txt"
    p.write_bytes("caf\xe9".encode("latin-1"))
    r = ex.extract_file(str(p), "a.txt")
    assert r.text_source == "text" and r.text == "café"
