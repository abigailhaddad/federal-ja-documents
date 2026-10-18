"""Spreadsheet builder: filtering, ordering, themes and workbook output. No network."""

import json
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from federal_ja.analysis import spreadsheet as ss

REV = "abc123"


def _notice(nid, posted, idv=False, matched=True):
    return {"notice_id": nid, "posted_date": posted, "agency": "AG", "sub_tier": "SUB", "office": "OFF",
            "title": f"T{nid}", "award_number": "0012AB", "sam_link": f"https://sam.gov/opp/{nid}/view",
            "usaspending_matched": matched, "usaspending_match_note": None, "usaspending_match_method": "exact",
            "usaspending_recipient_name": "VENDOR", "usaspending_recipient_uei": "0ABC",
            "usaspending_award_or_idv": "IDV" if idv else "AWARD", "usaspending_total_obligated": 0.0 if idv else 500.0,
            "usaspending_idv_orders_obligated": 9.0 if idv else None, "usaspending_potential_total_value": 1000.0,
            "usaspending_extent_competed": "NOT COMPETED", "usaspending_first_action_date": "2025-01-02",
            "usaspending_latest_action_date": "2025-02-03", "usaspending_permalink": "https://u", "usaspending_as_of": "2026-10-04"}


def _ext(nid, is_ja="yes", value="1234.5"):
    return {"notice_id": nid, "is_ja": is_ja, "vendor": "V", "estimated_value_usd": value, "summary_model_written": "S",
            "authority_plain": "only one responsible source", "authority_cited": "FAR 6.302-1", "doc_kind": "justification_and_approval",
            "value_basis": "total_including_options", "value_redacted": False, "value_quote": "q", "additional_vendors": None,
            "covers_multiple_actions": False, "period_of_performance": None, "signature_date": "2025-03-04",
            "market_research_summary_model_written": None, "other_firms_named": None, "redaction_detected": "yes",
            "values_dropped_by_checks": None, "notes_model_written": None, "truncated": False, "garbled_text_layer": False,
            "dataset_revision": REV, "run_at": "2026-10-18T00:00:00Z"}


def _status(nid, delay="tagged", reason=None):
    return {"notice_id": nid, "not_reviewed_reason": reason, "agency_delay_cause_status": delay,
            "sources_rejected_status": "not_tagged", "workforce_cause_status": "no_candidate_passages",
            "dataset_revision": REV}


def tables():
    notices = pd.DataFrame([_notice("a", "2025-05-01 10:00:00"), _notice("b", "2026-01-01 09:00:00", idv=True),
                            _notice("c", "2025-06-01 00:00:00"), _notice("d", "2025-07-01 00:00:00"),
                            _notice("e", "2025-08-01 00:00:00")])
    att = pd.DataFrame([{"notice_id": k, "attachment_order": 0, "file_path": f"files/{k}/x.pdf",
                         "download_status": "ok", "link_url": None} for k in "abcd"]
                       + [{"notice_id": "e", "attachment_order": 0, "file_path": None,
                           "download_status": "not_public", "link_url": None}])
    pages = pd.DataFrame([{"notice_id": k, "pages": 3, "transcribed": 1} for k in "abcd"])
    ext = pd.DataFrame([_ext("a"), _ext("b", value=None), _ext("c", is_ja="no"), _ext("d", is_ja="unclear")])
    status = pd.DataFrame([_status("a"), _status("b", reason="too_long"), _status("c"), _status("d")])
    tags = pd.DataFrame([{"notice_id": "a", "posted_date": "2025-05-01 10:00:00", "category": "agency_delay_cause",
                          "agency": "AG", "title": "Ta", "quote": "Q", "quote2": None, "explanation_model_written": "E",
                          "quote_pages": json.dumps([{"page_number": 2, "text_source": "transcription"}]),
                          "sam_link": "s"},
                         {"notice_id": "c", "posted_date": "2025-06-01 00:00:00", "category": "agency_delay_cause",
                          "agency": "AG", "title": "Tc", "quote": "Q", "quote2": None, "explanation_model_written": "E",
                          "quote_pages": "[]", "sam_link": "s"}])
    return {"notices": notices, "attachments": att, "pages": pages, "extractions": ext,
            "theme_status": status, "theme_tags": tags, "_data_rev": REV}


def test_filters_sorts_and_reconciles():
    df, tags, s = ss.build_frames(tables())
    assert list(df["Notice ID"]) == ["b", "a"]                      # only is_ja == yes, newest first
    assert s["left_out"] == {"no downloadable document": 1, "not a J&A (mostly award notices)": 1,
                             "unclear (blank templates, unreadable placeholders, garbled scans)": 1}
    assert s["rows"] + sum(s["left_out"].values()) == s["notices_total"]
    assert isinstance(df.loc[0, "Posted"], datetime)


def test_themes_and_vehicle_dollars():
    df, tags, _ = ss.build_frames(tables())
    a = df.set_index("Notice ID").loc["a"]
    b = df.set_index("Notice ID").loc["b"]
    assert a["Theme: agency's own delay caused it"] == "TAGGED"
    assert a["Theme: other firms responded but were rejected"] == "not tagged"
    assert a["Theme: shutdown / staff loss / DOGE / reorg caused it"] == "not tagged"
    assert b["Theme: agency's own delay caused it"] == "not reviewed (too long)"
    assert pd.isna(b["Obligated $ so far (USAspending)"])           # a contract vehicle has no money of its own
    assert a["Obligated $ so far (USAspending)"] == 500.0
    assert a["Estimated value $ stated in J&A (AI)"] == 1234.5
    assert a["Document"] == ss.HF_BASE + REV + "/files/a/x.pdf"      # pinned to the dataset commit
    assert list(tags["Notice ID"]) == ["a"]                          # tags on dropped documents are dropped
    assert tags.loc[0, "Quote page(s)"] == "p2 (transcribed)"


def test_workbook_ids_stay_text(tmp_path):
    df, tags, s = ss.build_frames(tables())
    out = tmp_path / "x.xlsx"
    ss.write_workbook(df, tags, ss.about_rows(s, "2026-10-18"), out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["About", "J&As", "Theme tags"]
    ws = wb["J&As"]
    hdr = [c.value for c in ws[1]]
    row = {h: c.value for h, c in zip(hdr, ws[2])}
    assert row["Award number"] == "0012AB" and row["Vendor UEI"] == "0ABC"   # leading zeros kept
    assert ws.freeze_panes == "C2"
    assert ws.cell(2, hdr.index("Document") + 1).hyperlink is not None
