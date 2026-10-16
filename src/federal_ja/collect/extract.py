"""Native text extraction and the "does this need transcription?" decision.

text_source values written to the dataset:
  pdf_text       text layer of a PDF, via pypdf
  docx           python-docx (paragraphs, then tables)
  xlsx           openpyxl (every sheet, tab-separated cells)
  text           plain-text attachment (.txt/.csv), decoded as UTF-8/latin-1
  transcription  machine transcription of page images (scanned PDFs)
  none           no text (unsupported type, unreadable file, or scanned PDF
                 not yet transcribed)

A PDF needs transcription when too few of its pages carry a usable text layer
(see ``pdf_needs_transcription``). Thresholds are module constants so they can
be audited and tuned.
"""

from __future__ import annotations

import io
import logging
import re
import zipfile
from dataclasses import dataclass, field

log = logging.getLogger(__name__)

# A page "has text" if its text layer has at least this many letters/digits.
# A typical typed J&A page has 1,500-4,000; a scanned page with no OCR has 0;
# a page with only a stamp or page number has well under 100.
MIN_ALNUM_PER_TEXT_PAGE = 100
# A PDF needs transcription if fewer than this share of pages has text.
MIN_TEXT_PAGE_SHARE = 0.5

_ALNUM = re.compile(r"[A-Za-z0-9]")


@dataclass
class ExtractResult:
    file_kind: str
    text: str = ""
    text_source: str = "none"
    page_count: int | None = None
    pages_with_text: int | None = None
    native_text_chars: int = 0
    needs_transcription: bool = False
    extract_error: str = ""
    page_texts: list[str] = field(default_factory=list, repr=False)


def detect_kind(data_head: bytes, filename: str = "") -> str:
    """Classify a file by magic bytes, falling back to its extension."""
    name = filename.lower()
    if data_head.startswith(b"%PDF") or b"%PDF-" in data_head[:1024]:
        return "pdf"
    if data_head.startswith(b"PK\x03\x04"):
        if name.endswith((".docx", ".docm")):
            return "docx"
        if name.endswith((".xlsx", ".xlsm")):
            return "xlsx"
        return "zip"  # resolved further in sniff_zip
    if data_head.startswith(b"\xd0\xcf\x11\xe0"):
        return "ole"  # legacy .doc/.xls
    if name.endswith((".txt", ".csv")):
        return "text"
    if name.endswith(".pdf"):
        return "pdf"  # damaged header; let pypdf try
    return "other"


def sniff_zip(path: str) -> str:
    try:
        with zipfile.ZipFile(path) as z:
            names = set(z.namelist())
    except zipfile.BadZipFile:
        return "other"
    if "word/document.xml" in names:
        return "docx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    return "zip"


def count_alnum(s: str) -> int:
    return len(_ALNUM.findall(s or ""))


def pdf_needs_transcription(page_texts: list[str]) -> bool:
    if not page_texts:
        return True
    with_text = sum(1 for t in page_texts if count_alnum(t) >= MIN_ALNUM_PER_TEXT_PAGE)
    return with_text / len(page_texts) < MIN_TEXT_PAGE_SHARE


def join_pages(page_texts: list[str]) -> str:
    return "\n\n".join(f"--- page {i} ---\n{t.strip()}" for i, t in enumerate(page_texts, 1))


def extract_pdf(path: str) -> ExtractResult:
    from pypdf import PdfReader

    res = ExtractResult(file_kind="pdf")
    try:
        reader = PdfReader(path)
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as e:  # noqa: BLE001
                res.extract_error = f"encrypted: {e}"
                return res
        texts = []
        for page in reader.pages:
            try:
                texts.append(page.extract_text() or "")
            except Exception as e:  # noqa: BLE001 - one bad page shouldn't sink the file
                log.warning("page extract failed in %s: %s", path, e)
                texts.append("")
    except Exception as e:  # noqa: BLE001
        res.extract_error = f"pypdf: {type(e).__name__}: {e}"[:500]
        return res
    res.page_texts = texts
    res.page_count = len(texts)
    res.pages_with_text = sum(1 for t in texts if count_alnum(t) >= MIN_ALNUM_PER_TEXT_PAGE)
    res.native_text_chars = sum(len(t) for t in texts)
    res.needs_transcription = pdf_needs_transcription(texts)
    if res.needs_transcription:
        # Keep whatever native text exists out of `text`; the transcription
        # will replace it. native_text_chars still records how much there was.
        res.text, res.text_source = "", "none"
    else:
        res.text, res.text_source = join_pages(texts), "pdf_text"
    return res


def extract_docx(path: str) -> ExtractResult:
    import docx

    res = ExtractResult(file_kind="docx")
    try:
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append("\t".join(c.text.strip() for c in row.cells))
        res.text = "\n".join(parts).strip()
    except Exception as e:  # noqa: BLE001
        res.extract_error = f"python-docx: {type(e).__name__}: {e}"[:500]
        return res
    res.native_text_chars = len(res.text)
    res.text_source = "docx" if res.text else "none"
    return res


def extract_xlsx(path: str) -> ExtractResult:
    import openpyxl

    res = ExtractResult(file_kind="xlsx")
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"=== sheet: {ws.title} ===")
            for row in ws.iter_rows(values_only=True):
                cells = ["" if v is None else str(v) for v in row]
                if any(cells):
                    parts.append("\t".join(cells).rstrip())
        wb.close()
        res.text = "\n".join(parts).strip()
    except Exception as e:  # noqa: BLE001
        res.extract_error = f"openpyxl: {type(e).__name__}: {e}"[:500]
        return res
    res.native_text_chars = len(res.text)
    res.text_source = "xlsx" if res.text else "none"
    return res


def extract_text_file(path: str) -> ExtractResult:
    res = ExtractResult(file_kind="text")
    raw = open(path, "rb").read()
    try:
        res.text = raw.decode("utf-8")
    except UnicodeDecodeError:
        res.text = raw.decode("latin-1")
    res.native_text_chars = len(res.text)
    res.text_source = "text" if res.text.strip() else "none"
    return res


def extract_file(path: str, filename: str = "") -> ExtractResult:
    with open(path, "rb") as f:
        head = f.read(2048)
    kind = detect_kind(head, filename)
    if kind == "zip":
        kind = sniff_zip(path)
    if kind == "pdf":
        return extract_pdf(path)
    if kind == "docx":
        return extract_docx(path)
    if kind == "xlsx":
        return extract_xlsx(path)
    if kind == "text":
        return extract_text_file(path)
    return ExtractResult(file_kind=kind)
